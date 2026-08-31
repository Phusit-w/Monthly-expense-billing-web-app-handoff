from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import time
import traceback
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

import psycopg
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pypdf import PdfReader

from ai_provider import ReviewRequest, build_provider


def load_local_env() -> None:
    env_file = Path(__file__).resolve().parent.parent / ".env"
    if not env_file.exists():
        return
    for raw in env_file.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


load_local_env()
_db_url = urlsplit(os.environ["DATABASE_URL"])
DATABASE_URL = urlunsplit((_db_url.scheme, _db_url.netloc, _db_url.path, "", _db_url.fragment))
DB_CONNECT_OPTIONS = (
    {
        "hostaddr": "127.0.0.1",
        "sslmode": "disable",
        "gssencmode": "disable",
        "connect_timeout": 10,
    }
    if _db_url.hostname in {"localhost", "127.0.0.1", "::1"}
    else {}
)
STORAGE_ROOT = Path(os.environ.get("SOC_STORAGE_ROOT", str(Path.cwd() / "data" / "soc"))).resolve()
POLL_SECONDS = int(os.environ.get("SOC_WORKER_POLL_SECONDS", "3"))
CHECK_LABELS = {
    "match": "ตรง", "mismatch": "ไม่ตรง", "review": "ต้องตรวจทาน",
    "not_found": "ไม่พบเลขหน้า", "unverifiable": "ยืนยันไม่ได้", "not_applicable": "ไม่ต้องตรวจ",
}
COLORS = {
    "match": "C6EFCE", "mismatch": "FFC7CE", "review": "FFEB9C",
    "not_found": "FFC7CE", "unverifiable": "FFEB9C", "not_applicable": "D9EAD3",
}


def db() -> psycopg.Connection:
    return psycopg.connect(
        DATABASE_URL,
        autocommit=False,
        prepare_threshold=None,
        **DB_CONNECT_OPTIONS,
    )


def safe_path(storage_key: str) -> Path:
    candidate = (STORAGE_ROOT / Path(storage_key)).resolve()
    if candidate != STORAGE_ROOT and STORAGE_ROOT not in candidate.parents:
        raise ValueError("invalid storage key")
    return candidate


def update_job(conn, job_id: str, **values) -> None:
    values["updatedAt"] = datetime.now(timezone.utc)
    assignments = ", ".join(f'"{key}" = %s' for key in values)
    conn.execute(f'UPDATE "SocJob" SET {assignments} WHERE id = %s', [*values.values(), job_id])
    conn.commit()


def claim_job(conn) -> tuple[str, str] | None:
    with conn.transaction():
        row = conn.execute(
            'SELECT id, status FROM "SocJob" WHERE status IN (\'QUEUED\', \'CONFIRMED\') AND "deletedAt" IS NULL '
            'ORDER BY "createdAt" FOR UPDATE SKIP LOCKED LIMIT 1'
        ).fetchone()
        if not row:
            return None
        next_status = "PROCESSING" if row[1] == "QUEUED" else "EXPORTING"
        stage = "กำลังอ่านเอกสาร" if next_status == "PROCESSING" else "กำลังสร้างเอกสาร Word"
        conn.execute(
            'UPDATE "SocJob" SET status=%s, stage=%s, progress=%s, "updatedAt"=NOW() WHERE id=%s',
            (next_status, stage, 10 if next_status == "PROCESSING" else 92, row[0]),
        )
        return row[0], next_status


def document_rows(conn, job_id: str, doc_type: str) -> list[dict]:
    rows = conn.execute(
        'SELECT id, "originalName", "storageKey", "mimeType" FROM "SocDocument" WHERE "jobId"=%s AND type=%s ORDER BY "createdAt"',
        (job_id, doc_type),
    ).fetchall()
    return [{"id": r[0], "name": r[1], "key": r[2], "mime": r[3]} for r in rows]


def normalized(value: str) -> str:
    return re.sub(r"[^\w\u0E00-\u0E7F]+", " ", value.lower(), flags=re.UNICODE).strip()


def unique_cell_texts(row) -> list[str]:
    values = [" ".join(cell.text.split()) for cell in row.cells]
    out: list[str] = []
    for value in values:
        if value and value not in out:
            out.append(value)
    return out


def header_score(values: list[str]) -> int:
    joined = normalized(" ".join(values))
    groups = [
        ("ลำดับ", "ข้อ", "item", "no"),
        ("tor", "ข้อกำหนด", "requirement", "specification"),
        ("อ้างอิง", "reference", "เลขหน้า", "page"),
    ]
    return sum(any(word in joined for word in group) for group in groups)


def is_item_label(value: str) -> bool:
    compact = re.sub(r"\s+", "", value)
    return bool(re.fullmatch(r"[0-9\u0E50-\u0E59]+(?:\.[0-9\u0E50-\u0E59]+)*\.?", compact))


def find_column(values: list[str], keywords: tuple[str, ...], fallback: int) -> int:
    for index, value in enumerate(values):
        text = normalized(value)
        if any(word in text for word in keywords):
            return index
    return min(fallback, max(0, len(values) - 1))


def classify_row(item: str, soc_text: str, cells: list[str]) -> str:
    text = normalized(soc_text)
    nonempty = [v for v in cells if v]
    if len(set(nonempty)) <= 1 or (not item and len(soc_text) < 90 and not re.search(r"\d", soc_text)):
        return "section_heading_row"
    if re.search(r"\b(model|รุ่น|ยี่ห้อ|brand|manufacturer|part\s*no)\b", text):
        return "product_identity_row"
    if len(soc_text) < 180 and re.search(r"(system|ระบบ|subsystem|module|equipment)\s*$", text):
        return "system_heading_row"
    return "content_row"


def parse_pages(reference: str) -> list[int]:
    matches = re.findall(r"(?:page(?:s)?|หน้า)\s*[:#.]?\s*([0-9,\-–—\s]+)", reference, flags=re.I)
    pages: set[int] = set()
    for match in matches:
        for part in re.split(r"[,\s]+", match.strip()):
            if not part:
                continue
            bounds = re.split(r"[-–—]", part)
            if len(bounds) == 2 and all(v.isdigit() for v in bounds):
                start, end = map(int, bounds)
                if 0 < start <= end <= start + 100:
                    pages.update(range(start, end + 1))
            elif part.isdigit() and int(part) > 0:
                pages.add(int(part))
    return sorted(pages)


def extract_soc_rows(path: Path) -> list[dict]:
    doc = Document(path)
    results: list[dict] = []
    physical_row = 0
    for table in doc.tables:
        if not table.rows:
            continue
        candidate = max(range(min(5, len(table.rows))), key=lambda i: header_score([c.text for c in table.rows[i].cells]))
        headers = [c.text for c in table.rows[candidate].cells]
        inferred_columns = header_score(headers) < 2
        if inferred_columns:
            data_rows = [
                index for index, row in enumerate(table.rows[:30])
                if len(row.cells) >= 3 and is_item_label(row.cells[0].text)
                and len(" ".join(row.cells[1].text.split())) >= 5
            ]
            if len(data_rows) < 3:
                continue
            candidate = data_rows[0] - 1
            item_col, tor_col, ref_col = 0, 1, len(table.columns) - 1
        else:
            item_col = find_column(headers, ("ลำดับ", "ข้อ", "item", "no"), 0)
            tor_col = find_column(headers, ("tor", "ข้อกำหนด", "requirement", "specification"), 1)
            ref_col = find_column(headers, ("อ้างอิง", "reference", "เลขหน้า", "page"), len(headers) - 1)
        for row in table.rows[candidate + 1:]:
            physical_row += 1
            raw = [" ".join(cell.text.split()) for cell in row.cells]
            if not any(raw):
                continue
            item = raw[item_col] if item_col < len(raw) else ""
            soc_text = raw[tor_col] if tor_col < len(raw) else ""
            reference = raw[ref_col] if ref_col < len(raw) else ""
            if inferred_columns and len(raw) > 2 and raw[2] and raw[2] != soc_text:
                soc_text = raw[2]
            if not soc_text:
                unique = unique_cell_texts(row)
                soc_text = max(unique, key=len, default="")
            if not soc_text:
                continue
            results.append({
                "row": physical_row, "item": item or str(physical_row), "soc_text": soc_text,
                "reference": reference, "pages": parse_pages(reference),
                "row_type": classify_row(item, soc_text, raw),
            })
    if not results:
        raise ValueError("ไม่พบตาราง SOC ที่มีคอลัมน์ข้อกำหนดและเลขอ้างอิง")
    return results


def pdf_text(path: Path) -> list[str]:
    reader = PdfReader(path)
    return [(page.extract_text() or "").strip() for page in reader.pages]


def choose_evidence(reference: str, evidence: list[dict]) -> dict:
    ref = normalized(reference)
    best = max(evidence, key=lambda d: len(set(normalized(d["name"]).split()) & set(ref.split())))
    return best


def meaningful_terms(value: str) -> set[str]:
    stop = {"the", "and", "with", "shall", "must", "หรือ", "และ", "ที่", "ให้", "มี", "ของ", "ต้อง"}
    return {t for t in normalized(value).split() if len(t) >= 3 and t not in stop}


def deterministic_review(row: dict, page_texts: list[str]) -> dict:
    row_type = row["row_type"]
    if row_type == "section_heading_row":
        return dict(reference="not_applicable", heading="not_applicable", product="not_applicable", relevance="not_applicable", detail="เป็นหัวข้อหมวดที่ไม่มี claim เฉพาะแถว", confidence="high")
    if not row["reference"].strip() or not row["pages"]:
        return dict(reference="not_found", heading="unverifiable" if row_type == "system_heading_row" else "not_applicable", product="unverifiable" if row_type == "product_identity_row" else "not_applicable", relevance="unverifiable", detail="ไม่พบเลขหน้าอ้างอิง จึงยังยืนยันหลักฐานไม่ได้", confidence="high")
    if any(not text or len(normalized(text)) < 30 for text in page_texts):
        return dict(reference="unverifiable", heading="unverifiable" if row_type == "system_heading_row" else "not_applicable", product="unverifiable" if row_type == "product_identity_row" else "not_applicable", relevance="unverifiable", detail="หน้าที่อ้างมีข้อความที่ดึงได้ไม่เพียงพอ ต้องตรวจจากภาพเอกสาร", confidence="low")
    joined = " ".join(page_texts)
    claim_terms = meaningful_terms(row["soc_text"])
    page_terms = meaningful_terms(joined)
    overlap = len(claim_terms & page_terms) / max(1, len(claim_terms))
    reference_status = "match" if overlap >= 0.45 else "review"
    heading = "not_applicable"
    if row_type == "system_heading_row":
        heading = "match" if normalized(row["soc_text"]) in normalized(joined) else "unverifiable"
    product = "not_applicable"
    if row_type == "product_identity_row":
        product = "match" if overlap >= 0.6 else "unverifiable"
    relevance = "related" if overlap >= 0.25 else "review"
    detail = f"ตรวจหน้า {', '.join(map(str, row['pages']))}; พบคำสำคัญร่วมกันประมาณ {round(overlap * 100)}%"
    if reference_status == "review":
        detail += " จึงต้องตรวจบริบทบนหน้าจริงก่อนสรุป"
    return dict(reference=reference_status, heading=heading, product=product, relevance=relevance, detail=detail, confidence="medium" if reference_status == "match" else "low")


def process_audit(conn, job_id: str) -> None:
    soc_docs = document_rows(conn, job_id, "SOC")
    evidence_docs = document_rows(conn, job_id, "EVIDENCE")
    if len(soc_docs) != 1 or not evidence_docs:
        raise ValueError("ไฟล์ประกอบงานไม่ครบ")
    update_job(conn, job_id, stage="กำลังแยกตาราง SOC", progress=20)
    rows = extract_soc_rows(safe_path(soc_docs[0]["key"]))
    update_job(conn, job_id, stage=f"พบ {len(rows)} รายการ กำลังอ่านข้อความ PDF", progress=35)
    evidence_text = {doc["id"]: pdf_text(safe_path(doc["key"])) for doc in evidence_docs}
    provider = build_provider()
    payloads = []
    for index, row in enumerate(rows, start=1):
        evidence = choose_evidence(row["reference"], evidence_docs)
        pages = evidence_text[evidence["id"]]
        out_of_range = [p for p in row["pages"] if p > len(pages)]
        if out_of_range:
            review = dict(reference="mismatch", heading="unverifiable" if row["row_type"] == "system_heading_row" else "not_applicable", product="unverifiable" if row["row_type"] == "product_identity_row" else "not_applicable", relevance="unverifiable", detail=f"เลขหน้า {', '.join(map(str, out_of_range))} เกินจำนวนหน้า PDF ({len(pages)} หน้า)", confidence="high")
            selected_texts: list[str] = []
        else:
            selected_texts = [pages[p - 1] for p in row["pages"]]
            review = deterministic_review(row, selected_texts)
            ai = provider.review(ReviewRequest(row["row_type"], row["soc_text"], row["reference"], selected_texts))
            if ai:
                review.update(ai)
        payloads.append((
            uuid.uuid4().hex, job_id, row["row"], row["item"][:200], row["row_type"], row["soc_text"],
            row["reference"], json.dumps(row["pages"]), evidence["id"], review["reference"], review["heading"],
            review["product"], review["relevance"], review["detail"], review["confidence"], review["reference"],
            review["heading"], review["product"], review["relevance"], review["detail"],
        ))
        if index % 10 == 0:
            update_job(conn, job_id, stage=f"ตรวจแล้ว {index}/{len(rows)} รายการ", progress=min(80, 40 + int(index / len(rows) * 40)))
    with conn.transaction():
        conn.execute('DELETE FROM "SocCheckResult" WHERE "jobId"=%s', (job_id,))
        with conn.cursor() as cursor:
            statement = (
                'INSERT INTO "SocCheckResult" (id,"jobId","rowNumber",item,"rowType","socText","referenceText","referencePages","evidenceDocumentId","aiReferenceCheck","aiHeadingTitleCheck","aiProductIdentity","aiContentRelevance","aiDetail","aiConfidence","finalReferenceCheck","finalHeadingTitleCheck","finalProductIdentity","finalContentRelevance","finalDetail","createdAt","updatedAt") '
                'VALUES (%s,%s,%s,%s,%s,%s,%s,%s::jsonb,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,NOW(),NOW())'
            )
            for payload in payloads:
                cursor.execute(statement, payload)
        conn.execute('UPDATE "SocJob" SET status=\'NEEDS_REVIEW\', stage=\'รอผู้ใช้ตรวจทาน\', progress=85, "updatedAt"=NOW() WHERE id=%s', (job_id,))
        conn.execute('INSERT INTO "SocAuditEvent" (id,"jobId",action,detail,"createdAt") VALUES (%s,%s,\'AUDIT_COMPLETED\',%s::jsonb,NOW())', (uuid.uuid4().hex, job_id, json.dumps({"resultCount": len(payloads)})))


def set_cell_text(cell, value: object, bold: bool = False, size: float = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value or ""))
    run.bold = bold
    run.font.name = "TH Sarabun New"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)


def fill(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), color)


def append_report(input_path: Path, output_path: Path, title: str, results: list[dict]) -> None:
    doc = Document(input_path)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Cm(1.2)
    section.left_margin = section.right_margin = Cm(1.2)
    section.header.is_linked_to_previous = False
    for table in list(section.header.tables):
        table._element.getparent().remove(table._element)
    for paragraph in list(section.header.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)
    section.header.add_paragraph("")
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run = heading.add_run(f"ผลการตรวจสอบการอ้างอิง SOC — {title}")
    set_run.bold = True
    set_run.font.name = "TH Sarabun New"
    set_run.font.size = Pt(18)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"วันที่สร้างผล {datetime.now().strftime('%d/%m/%Y')} | ผ่านการยืนยันโดยผู้ใช้งาน")
    counts = Counter(r["reference"] for r in results)
    summary = doc.add_table(rows=2, cols=6)
    summary.style = "Table Grid"
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    order = ["match", "mismatch", "review", "unverifiable", "not_found", "not_applicable"]
    for i, status in enumerate(order):
        set_cell_text(summary.cell(0, i), CHECK_LABELS[status], True)
        fill(summary.cell(0, i), COLORS[status])
        set_cell_text(summary.cell(1, i), counts.get(status, 0), True, 10)
        summary.cell(1, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["แถว", "ข้อ", "เลขหน้าอ้างอิง", "ผลอ้างอิง", "ผลชื่อหัวข้อ", "ความมั่นใจ", "รายละเอียด"]
    for i, label in enumerate(headers):
        set_cell_text(table.cell(0, i), label, True)
        fill(table.cell(0, i), "D9EAF7")
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(marker)
    for result in results:
        cells = table.add_row().cells
        values = [result["row"], result["item"], result["reference_text"] or "—", CHECK_LABELS.get(result["reference"], result["reference"]), CHECK_LABELS.get(result["heading"], result["heading"]), result["confidence"], result["detail"]]
        for cell, value in zip(cells, values):
            set_cell_text(cell, value, size=7.5)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        fill(cells[3], COLORS.get(result["reference"], "FFFFFF"))
        fill(cells[4], COLORS.get(result["heading"], "FFFFFF"))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def process_export(conn, job_id: str) -> None:
    soc = document_rows(conn, job_id, "SOC")[0]
    title = conn.execute('SELECT title FROM "SocJob" WHERE id=%s', (job_id,)).fetchone()[0]
    rows = conn.execute(
        'SELECT "rowNumber",item,"referenceText","finalReferenceCheck","finalHeadingTitleCheck","aiConfidence","finalDetail" FROM "SocCheckResult" WHERE "jobId"=%s ORDER BY "rowNumber"',
        (job_id,),
    ).fetchall()
    results = [{"row": r[0], "item": r[1], "reference_text": r[2], "reference": r[3], "heading": r[4], "confidence": r[5], "detail": r[6]} for r in rows]
    if not results:
        raise ValueError("ไม่พบผลตรวจสำหรับสร้างเอกสาร")
    slug = re.sub(r"[^\w\u0E00-\u0E7F-]+", "-", title).strip("-")[:80] or "SOC"
    filename = f"SOC_Check-{datetime.now().strftime('%Y-%m-%d')}-{slug}.docx"
    key = f"{job_id}/{uuid.uuid4().hex}.docx"
    output = safe_path(key)
    append_report(safe_path(soc["key"]), output, title, results)
    checksum = hashlib.sha256(output.read_bytes()).hexdigest()
    preview_key = None
    libreoffice = shutil.which("libreoffice")
    if libreoffice:
        preview_dir = STORAGE_ROOT / job_id / "preview"
        preview_dir.mkdir(parents=True, exist_ok=True)
        subprocess.run([libreoffice, "--headless", "--convert-to", "pdf", "--outdir", str(preview_dir), str(output)], check=False, timeout=120)
        generated = preview_dir / f"{output.stem}.pdf"
        if generated.exists():
            preview_key = str(generated.relative_to(STORAGE_ROOT)).replace("\\", "/")
    with conn.transaction():
        conn.execute('DELETE FROM "SocDocument" WHERE "jobId"=%s AND type IN (\'OUTPUT\',\'PREVIEW\')', (job_id,))
        conn.execute(
            'INSERT INTO "SocDocument" (id,"jobId",type,"originalName","storageKey","mimeType","sizeBytes",checksum,"createdAt") VALUES (%s,%s,\'OUTPUT\',%s,%s,%s,%s,%s,NOW())',
            (uuid.uuid4().hex, job_id, filename, key, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", output.stat().st_size, checksum),
        )
        if preview_key:
            preview = safe_path(preview_key)
            conn.execute(
                'INSERT INTO "SocDocument" (id,"jobId",type,"originalName","storageKey","mimeType","sizeBytes",checksum,"createdAt") VALUES (%s,%s,\'PREVIEW\',%s,%s,\'application/pdf\',%s,%s,NOW())',
                (uuid.uuid4().hex, job_id, f"{slug}-preview.pdf", preview_key, preview.stat().st_size, hashlib.sha256(preview.read_bytes()).hexdigest()),
            )
        conn.execute('UPDATE "SocJob" SET status=\'COMPLETED\', stage=\'เสร็จสิ้น\', progress=100, "completedAt"=NOW(), "updatedAt"=NOW() WHERE id=%s', (job_id,))
        conn.execute('INSERT INTO "SocAuditEvent" (id,"jobId",action,"createdAt") VALUES (%s,%s,\'EXPORT_COMPLETED\',NOW())', (uuid.uuid4().hex, job_id))


def cleanup_expired(conn) -> None:
    rows = conn.execute('SELECT id FROM "SocJob" WHERE "expiresAt" < NOW() AND status <> \'EXPIRED\' LIMIT 20').fetchall()
    for (job_id,) in rows:
        shutil.rmtree(STORAGE_ROOT / job_id, ignore_errors=True)
        with conn.transaction():
            conn.execute('DELETE FROM "SocDocument" WHERE "jobId"=%s', (job_id,))
            conn.execute('UPDATE "SocCheckResult" SET "socText"=\'[หมดอายุ]\', "referenceText"=\'\', "aiDetail"=\'[หมดอายุ]\', "finalDetail"=\'[หมดอายุ]\', "referencePages"=\'[]\'::jsonb WHERE "jobId"=%s', (job_id,))
            conn.execute('UPDATE "SocJob" SET status=\'EXPIRED\', stage=\'ลบไฟล์ตามนโยบาย 90 วันแล้ว\', progress=100, "updatedAt"=NOW() WHERE id=%s', (job_id,))


def maintain_system(conn) -> None:
    conn.execute(
        'INSERT INTO "WorkerHeartbeat" (name,status,"lastSeenAt","updatedAt") VALUES (\'soc-worker\',\'RUNNING\',NOW(),NOW()) '
        'ON CONFLICT (name) DO UPDATE SET status=\'RUNNING\', "lastSeenAt"=NOW(), "updatedAt"=NOW(), "lastError"=NULL'
    )
    rows = conn.execute('SELECT id FROM "SocJob" WHERE "purgeAfter" < NOW() AND "purgedAt" IS NULL LIMIT 20').fetchall()
    for (job_id,) in rows:
        shutil.rmtree(STORAGE_ROOT / job_id, ignore_errors=True)
        with conn.transaction():
            conn.execute('DELETE FROM "SocDocument" WHERE "jobId"=%s', (job_id,))
            conn.execute('UPDATE "SocCheckResult" SET "socText"=\'[ล้างข้อมูลแล้ว]\', "referenceText"=\'\', "aiDetail"=\'[ล้างข้อมูลแล้ว]\', "finalDetail"=\'[ล้างข้อมูลแล้ว]\', "referencePages"=\'[]\'::jsonb WHERE "jobId"=%s', (job_id,))
            conn.execute('UPDATE "SocJob" SET "purgedAt"=NOW(), stage=\'ล้างข้อมูลจากถังขยะแล้ว\', "updatedAt"=NOW() WHERE id=%s', (job_id,))
    conn.execute('DELETE FROM "ExpenseRecord" WHERE "purgeAfter" < NOW() AND "deletedAt" IS NOT NULL')
    conn.execute('DELETE FROM "AuditLog" WHERE "createdAt" < NOW() - INTERVAL \'730 days\'')
    conn.commit()


def main() -> None:
    STORAGE_ROOT.mkdir(parents=True, exist_ok=True)
    while True:
        try:
            with db() as conn:
                while True:
                    maintain_system(conn)
                    cleanup_expired(conn)
                    claimed = claim_job(conn)
                    if not claimed:
                        time.sleep(POLL_SECONDS)
                        continue
                    job_id, status = claimed
                    try:
                        if status == "PROCESSING":
                            process_audit(conn, job_id)
                        else:
                            process_export(conn, job_id)
                    except Exception as exc:
                        traceback.print_exc()
                        update_job(conn, job_id, status="FAILED", stage="เกิดข้อผิดพลาด", errorMessage=str(exc)[:1000])
        except Exception:
            traceback.print_exc()
            time.sleep(POLL_SECONDS)


if __name__ == "__main__":
    main()
