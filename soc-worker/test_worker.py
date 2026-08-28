from __future__ import annotations

import os
import sys
import unittest
from pathlib import Path

os.environ.setdefault("DATABASE_URL", "postgresql://unused:unused@localhost/unused")
sys.path.insert(0, str(Path(__file__).parent))

from docx import Document

from worker import append_report, deterministic_review, extract_soc_rows, parse_pages


class WorkerRulesTest(unittest.TestCase):
    def setUp(self):
        self.output_dir = Path(__file__).parent / "test-output"
        self.output_dir.mkdir(exist_ok=True)

    def tearDown(self):
        for path in self.output_dir.iterdir():
            if path.is_file():
                path.unlink()

    def test_parse_pages_supports_lists_and_ranges(self):
        self.assertEqual(parse_pages("Catalog page 3, 5-7"), [3, 5, 6, 7])
        self.assertEqual(parse_pages("หน้า 12–13"), [12, 13])

    def test_extracts_dynamic_columns_and_classifies_rows(self):
        source = self.output_dir / "soc.docx"
        try:
            doc = Document()
            table = doc.add_table(rows=1, cols=4)
            for cell, value in zip(table.rows[0].cells, ["ข้อ", "รายละเอียดข้อกำหนด TOR", "Comply", "เลขอ้างอิงในเอกสาร"]):
                cell.text = value
            row = table.add_row().cells
            for cell, value in zip(row, ["3.2", "ระบบต้องรองรับการจัดการอุปกรณ์", "✓", "Catalog, page 3"]):
                cell.text = value
            doc.save(source)
            rows = extract_soc_rows(source)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["item"], "3.2")
            self.assertEqual(rows[0]["pages"], [3])
            self.assertEqual(rows[0]["row_type"], "content_row")
        finally:
            source.unlink(missing_ok=True)

    def test_extracts_headerless_appendix_table(self):
        source = self.output_dir / "headerless-soc.docx"
        try:
            doc = Document()
            table = doc.add_table(rows=1, cols=7)
            table.rows[0].cells[0].text = "ภาคผนวก ก. คุณลักษณะเฉพาะ"
            for item, requirement, proposal, reference in [
                ("๑.๑.๑", "ความละเอียดไม่น้อยกว่า 70 ล้านพิกเซล", "116 million pixel", "Catalog, page 2"),
                ("๑.๑.๒", "มุมมองแนวนอนไม่น้อยกว่า 35 องศา", "horizontal field of view 180 degrees", "Catalog, page 3"),
                ("๑.๑.๓", "เซนเซอร์ CMOS ไม่น้อยกว่า 8 ล้านพิกเซล", "CMOS resolution 8.3 million", "Catalog, page 4"),
            ]:
                cells = table.add_row().cells
                cells[0].text, cells[1].text, cells[2].text, cells[6].text = item, requirement, proposal, reference
            doc.save(source)
            rows = extract_soc_rows(source)
            self.assertEqual(len(rows), 3)
            self.assertEqual(rows[0]["item"], "๑.๑.๑")
            self.assertEqual(rows[0]["soc_text"], "116 million pixel")
            self.assertEqual(rows[0]["pages"], [2])
        finally:
            source.unlink(missing_ok=True)

    def test_unreadable_page_never_becomes_mismatch(self):
        result = deterministic_review({
            "row_type": "content_row", "reference": "page 1", "pages": [1],
            "soc_text": "รองรับการบริหารจัดการระบบ",
        }, [""])
        self.assertEqual(result["reference"], "unverifiable")

    def test_export_appends_report_without_changing_source(self):
        source = self.output_dir / "source.docx"
        output = self.output_dir / "output.docx"
        try:
            doc = Document()
            doc.add_paragraph("SOC ORIGINAL")
            doc.save(source)
            before = source.read_bytes()
            append_report(source, output, "Demo", [{
                "row": 1, "item": "1.1", "reference_text": "page 1", "reference": "match",
                "heading": "not_applicable", "confidence": "high", "detail": "พบหลักฐานหน้า 1",
            }])
            self.assertEqual(source.read_bytes(), before)
            result = Document(output)
            self.assertIn("SOC ORIGINAL", "\n".join(p.text for p in result.paragraphs))
            self.assertGreaterEqual(len(result.tables), 2)
        finally:
            source.unlink(missing_ok=True)
            output.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
