from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from docx import Document

from extract_text import MIN_USABLE_TEXT_LENGTH, extract_text, find_source_document


class FindSourceDocumentTest(unittest.TestCase):
    def test_prefers_docx_over_pdf_in_project_root(self):
        with tempfile.TemporaryDirectory() as tmp:
            project = Path(tmp)
            (project / "notes.pdf").write_bytes(b"%PDF-1.4 fake")
            (project / "TOR.docx").write_bytes(b"fake docx bytes")
            self.assertEqual(find_source_document(project).name, "TOR.docx")

    def test_looks_inside_tor_and_proposal_subfolders(self):
        with tempfile.TemporaryDirectory() as tmp:
            project = Path(tmp)
            (project / "_TOR").mkdir()
            (project / "_TOR" / "tor_document.docx").write_bytes(b"fake docx bytes")
            self.assertEqual(find_source_document(project).name, "tor_document.docx")

    def test_returns_none_when_nothing_usable_is_present(self):
        with tempfile.TemporaryDirectory() as tmp:
            project = Path(tmp)
            (project / "photo.jpg").write_bytes(b"not a document")
            self.assertIsNone(find_source_document(project))


class ExtractTextTest(unittest.TestCase):
    def test_extracts_docx_paragraph_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "proposal.docx"
            doc = Document()
            # Padded well past MIN_USABLE_TEXT_LENGTH so this exercises the
            # real "usable text" branch, not the too-short fallback.
            doc.add_paragraph("โครงการติดตั้งระบบโซลาร์รูฟท็อป " * 20)
            doc.save(path)
            text = extract_text(path)
            self.assertIsNotNone(text)
            self.assertIn("โซลาร์รูฟท็อป", text)

    def test_short_docx_is_treated_as_unusable(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.docx"
            doc = Document()
            doc.add_paragraph("สั้นเกินไป")
            doc.save(path)
            self.assertLess(len("สั้นเกินไป"), MIN_USABLE_TEXT_LENGTH)
            self.assertIsNone(extract_text(path))

    def test_unreadable_file_returns_none_instead_of_raising(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "corrupt.docx"
            path.write_bytes(b"not actually a docx file")
            self.assertIsNone(extract_text(path))

    def test_unsupported_extension_returns_none(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "notes.txt"
            path.write_text("x" * (MIN_USABLE_TEXT_LENGTH + 10))
            self.assertIsNone(extract_text(path))


if __name__ == "__main__":
    unittest.main()
