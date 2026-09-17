from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from docx import Document

from ai_provider import DisabledProvider, SummarizeResult
from crawler import build_payload, crawl
from discovery import ProjectFolder


class FakeProvider:
    def __init__(self, result: SummarizeResult | None):
        self._result = result

    def summarize(self, request):
        return self._result


class BuildPayloadTest(unittest.TestCase):
    def test_disabled_provider_leaves_description_and_budget_blank(self):
        with tempfile.TemporaryDirectory() as tmp:
            project_dir = Path(tmp) / "MOFXXX_RFID"
            project_dir.mkdir()
            project = ProjectFolder(client="MOF", path=project_dir, year_hint=2026)

            payload = build_payload(project, DisabledProvider())

            self.assertEqual(payload["folderPath"], str(project_dir))
            self.assertEqual(payload["client"], "MOF")
            self.assertEqual(payload["projectName"], "MOFXXX_RFID")
            self.assertEqual(payload["descriptionTh"], "")
            self.assertEqual(payload["descriptionEn"], "")
            self.assertIsNone(payload["budgetAmount"])
            self.assertEqual(payload["year"], 2026)  # structural hint, no AI needed

    def test_structural_year_hint_wins_over_ai_guessed_year(self):
        with tempfile.TemporaryDirectory() as tmp:
            project_dir = Path(tmp) / "NTXXX_Fiber"
            project_dir.mkdir()
            doc = Document()
            doc.add_paragraph("โครงการใยแก้วนำแสงสำหรับพื้นที่ชนบท " * 20)
            doc.save(project_dir / "TOR.docx")
            project = ProjectFolder(client="NT", path=project_dir, year_hint=2024)
            provider = FakeProvider(SummarizeResult(description_th="a", description_en="b", budget_amount=100.0, year=1999))

            payload = build_payload(project, provider)

            self.assertEqual(payload["year"], 2024)

    def test_ai_year_used_only_when_no_structural_hint(self):
        with tempfile.TemporaryDirectory() as tmp:
            project_dir = Path(tmp) / "MOFYYY_OldOne"
            project_dir.mkdir()
            doc = Document()
            doc.add_paragraph("รายละเอียดโครงการเก่าที่ไม่มีปีระบุในโครงสร้างโฟลเดอร์ " * 20)
            doc.save(project_dir / "TOR.docx")
            project = ProjectFolder(client="MOF", path=project_dir, year_hint=None)
            provider = FakeProvider(SummarizeResult(description_th="a", description_en="b", budget_amount=100.0, year=2019))

            payload = build_payload(project, provider)

            self.assertEqual(payload["year"], 2019)
            self.assertEqual(payload["budgetAmount"], 100.0)

    def test_no_source_document_leaves_everything_blank_without_calling_the_provider(self):
        with tempfile.TemporaryDirectory() as tmp:
            project_dir = Path(tmp) / "NoDocs"
            project_dir.mkdir()
            project = ProjectFolder(client="MOF", path=project_dir, year_hint=None)

            def fail_if_called(request):
                raise AssertionError("summarize() should not be called when there's no extractable text")

            provider = FakeProvider(None)
            provider.summarize = fail_if_called  # type: ignore[method-assign]

            payload = build_payload(project, provider)
            self.assertEqual(payload["descriptionTh"], "")
            self.assertIsNone(payload["year"])


class CrawlTest(unittest.TestCase):
    def test_crawl_reports_stats_and_skipped_folders(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "_Project 2026/MOF/MOFXXX_RFID/_TOR").mkdir(parents=True)
            (root / "_Project 2026/MOF/unrecognised_folder/random").mkdir(parents=True)

            payloads, stats = crawl(root, DisabledProvider())

            self.assertEqual(stats.total_projects, 1)
            self.assertEqual(stats.summarized, 0)
            self.assertEqual(len(stats.skipped_folders), 1)
            self.assertEqual(len(payloads), 1)
            self.assertEqual(payloads[0]["projectName"], "MOFXXX_RFID")


if __name__ == "__main__":
    unittest.main()
